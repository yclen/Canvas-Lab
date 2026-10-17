"""fit_spectral_response.py — count rate vs. power, fitted with a line, for every run in a folder.

    python scripts/fit_spectral_response.py                    # asks for the folder
    python scripts/fit_spectral_response.py "experiments/detector_sweep/Folder 1"

For a folder of detector_spectral_response.py runs (sweep_<wavelength>nm_<time>.csv):

- fits each run's count rate vs. power with a straight line,
      count_rate = slope * power_mW + intercept
  (the slope is the detector's response at that wavelength, in counts/s per mW;
  the intercept is about the dark count rate). Points the run marked clipped or
  overflowed are left out of the fit (drawn hollow).
- prints each run's wavelength, slope, intercept and R²,
- plots every run's data and its fit line on one plot (shown, and saved as
  spectral_response_fits_<time>.png in the folder),
- plots each run's slope vs. its wavelength, the detector's response (shown,
  and saved as spectral_response_fits_<time>_slopes.png),
- writes spectral_response_fits_<time>.xlsx in the folder: one row per run —
  wavelength, slope, intercept, R², points, the equation y = mx + b, the file.

The wavelength comes from the run's .xlsx ("settings": wavelength_nm), else its
file name. Runs at the same wavelength are each listed (by their start time).
"""

import argparse
import csv
import glob
import os
import re
import sys
from datetime import datetime

import numpy as np

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DEFAULT_DIR = os.path.join(ROOT, "experiments", "detector_sweep")
X, Y = "power_mW", "count_rate"
NAME = re.compile(r"_(\d+(?:\.\d+)?)nm_(\d{8}_\d{6})")


def pick_folder():
    """A folder chooser (starting in experiments/detector_sweep/), or None if cancelled."""
    import tkinter
    from tkinter import filedialog
    root = tkinter.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    folder = filedialog.askdirectory(title="Pick a folder of detector spectral response runs",
                                     initialdir=DEFAULT_DIR if os.path.isdir(DEFAULT_DIR) else ROOT)
    root.destroy()
    return folder or None


def wavelength_of(csv_path):
    """The run's wavelength (nm): its .xlsx settings, else the file name; None if neither says."""
    xlsx = os.path.splitext(csv_path)[0] + ".xlsx"
    if os.path.exists(xlsx):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx, read_only=True)
            if "settings" in wb.sheetnames:
                for key, value in wb["settings"].iter_rows(values_only=True):
                    if key == "wavelength_nm" and value is not None:
                        return float(value)
        except Exception:
            pass                                            # unreadable: fall back to the name
    m = NAME.search(os.path.basename(csv_path))
    return float(m.group(1)) if m else None


def read_run(csv_path):
    """(power, count rate, used-in-fit mask) as arrays, or None if it has no such columns."""
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if not rows or X not in rows[0] or Y not in rows[0]:
        return None
    x, y, ok = [], [], []
    for r in rows:
        try:
            px, py = float(r[X]), float(r[Y])
        except (TypeError, ValueError):
            continue
        if not (np.isfinite(px) and np.isfinite(py)):
            continue
        bad = any(str(r.get(k, "0")).strip() not in ("", "0", "0.0", "False") for k in ("clipped", "overflows"))
        x.append(px)
        y.append(py)
        ok.append(not bad)
    return np.array(x), np.array(y), np.array(ok, dtype=bool)


def fit_line(x, y):
    """(slope, intercept, R²) of the least-squares line y = slope * x + intercept."""
    slope, intercept = np.polyfit(x, y, 1)
    residual = y - (slope * x + intercept)
    total = np.sum((y - y.mean()) ** 2)
    r2 = 1 - np.sum(residual ** 2) / total if total > 0 else float("nan")
    return float(slope), float(intercept), float(r2)


def equation(slope, intercept):
    sign = "-" if intercept < 0 else "+"
    return f"y = {slope:.6g} x {sign} {abs(intercept):.6g}"


def main():
    p = argparse.ArgumentParser(description="Fit count rate vs. power for every run in a folder.")
    p.add_argument("folder", nargs="?", help="the folder of runs (default: asks)")
    p.add_argument("--no-show", action="store_true", help="save the plot without opening a window")
    args = p.parse_args()
    folder = args.folder or pick_folder()
    if not folder:
        print("No folder picked.")
        return 1
    folder = os.path.abspath(folder)

    runs = []
    for path in sorted(glob.glob(os.path.join(folder, "*.csv"))):
        data = read_run(path)
        if data is None:
            continue
        x, y, ok = data
        if ok.sum() < 2:
            print(f"skipped {os.path.basename(path)}: fewer than 2 usable points")
            continue
        m = NAME.search(os.path.basename(path))
        started = datetime.strptime(m.group(2), "%Y%m%d_%H%M%S") if m else None
        runs.append({"file": os.path.basename(path), "wavelength": wavelength_of(path), "started": started,
                     "x": x, "y": y, "ok": ok, "fit": fit_line(x[ok], y[ok])})
    if not runs:
        print(f"No runs with {X} and {Y} columns in {folder}")
        return 1
    runs.sort(key=lambda r: (r["wavelength"] is None, r["wavelength"] or 0, r["started"] or datetime.min))

    # The same wavelength more than once: tell them apart by start time.
    counts = {}
    for r in runs:
        counts[r["wavelength"]] = counts.get(r["wavelength"], 0) + 1
    for r in runs:
        name = "? nm" if r["wavelength"] is None else f"{r['wavelength']:g} nm"
        if counts[r["wavelength"]] > 1 and r["started"]:
            name += f" ({r['started']:%H:%M})"
        r["label"] = name

    # The list.
    print(f"\n{folder}\n")
    print(f"{'run':<18} {'slope (counts/s per mW)':>24} {'intercept (counts/s)':>21} {'R^2':>8} {'points':>7}")
    for r in runs:
        slope, intercept, r2 = r["fit"]
        print(f"{r['label']:<18} {slope:>24.6g} {intercept:>21.6g} {r2:>8.5f} {int(r['ok'].sum()):>7}")
    print()
    for r in runs:
        print(f"{r['label']:<18} {equation(*r['fit'][:2])}")

    # The plot: every run's points and its line.
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = os.path.join(folder, f"spectral_response_fits_{stamp}")
    import matplotlib
    if args.no_show:
        matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(9, 6))
    colors = plt.cm.viridis(np.linspace(0, 0.9, len(runs)))
    for r, color in zip(runs, colors):
        x, y, ok = r["x"], r["y"], r["ok"]
        slope, intercept, r2 = r["fit"]
        ax.plot(x[ok], y[ok], "o", color=color, ms=5, label=f"{r['label']}  m={slope:.4g}")
        if (~ok).any():
            ax.plot(x[~ok], y[~ok], "o", mfc="none", color=color, ms=5)
        line = np.linspace(0, x.max(), 50)
        ax.plot(line, slope * line + intercept, "-", color=color, lw=1.2)
    ax.set_xlabel("power (mW)")
    ax.set_ylabel("count rate (counts/s)")
    ax.set_title(f"Count rate vs. power, linear fits — {os.path.basename(folder)}")
    ax.grid(alpha=0.3)
    ax.legend(fontsize=8, title="run  (slope, counts/s per mW)", title_fontsize=8)
    fig.tight_layout()
    fig.savefig(base + ".png", dpi=150)

    # The response: each run's slope vs. its wavelength (runs at the same wavelength: each a
    # point, the line through their mean).
    known = [r for r in runs if r["wavelength"] is not None]
    if known:
        fig2, ax2 = plt.subplots(figsize=(7, 4.5))
        ax2.plot([r["wavelength"] for r in known], [r["fit"][0] for r in known], "o", color="tab:blue", ms=6)
        waves = sorted({r["wavelength"] for r in known})
        means = [np.mean([r["fit"][0] for r in known if r["wavelength"] == w]) for w in waves]
        ax2.plot(waves, means, "-", color="tab:blue", lw=1.2, alpha=0.7)
        ax2.set_xlabel("wavelength (nm)")
        ax2.set_ylabel("slope (counts/s per mW)")
        ax2.set_title(f"Detector response vs. wavelength — {os.path.basename(folder)}")
        ax2.set_ylim(bottom=0)
        ax2.grid(alpha=0.3)
        fig2.tight_layout()
        fig2.savefig(base + "_slopes.png", dpi=150)

    # The table.
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "fits"
    header = ["wavelength (nm)", "slope m (counts/s per mW)", "intercept b (counts/s)", "R²",
              "points fitted", "equation (y = count rate, x = power in mW)", "started", "file"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in runs:
        slope, intercept, r2 = r["fit"]
        ws.append([r["wavelength"], slope, intercept, r2, int(r["ok"].sum()), equation(slope, intercept),
                   r["started"], r["file"]])
    for col, width in zip("ABCDEFGH", (16, 26, 22, 10, 13, 42, 19, 36)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[6].number_format = "yyyy-mm-dd hh:mm:ss"
    ws.freeze_panes = "A2"
    wb.save(base + ".xlsx")

    print(f"\nwrote {base}.xlsx\n      {base}.png" + (f"\n      {base}_slopes.png" if known else ""))
    if not args.no_show:
        plt.show()
    return 0


if __name__ == "__main__":
    sys.exit(main())
