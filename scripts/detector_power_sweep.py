"""detector_power_sweep.py — PDA45 counts and voltage vs. excitation power, at one wavelength.

    python scripts/detector_power_sweep.py --wavelength 550
    python scripts/detector_power_sweep.py --wavelength 600 --width 10 --start 90 --stop 60 --points 20
    python scripts/detector_power_sweep.py --wavelength 550 --spacing power    # evenly spaced in mW

1. Sets the excitation band to --wavelength (center) and --width (nm), and the
   power meter's correction wavelength to match.
2. Power pass: flips the detection mirror UP (beam -> PM100D) and, at each of
   --points ND stage positions from --start to --stop, averages a few power
   readings.
3. Detector pass: flips the mirror DOWN (beam -> sphere -> PDA45) and goes
   back through the same positions, averaging the PDA45's count rate (Time
   Tagger) and mean voltage (NI DAQ).

So each position gets a power, a count rate and a voltage. The run is written
to experiments/detector_sweep/ as <name>.xlsx (sheet "data": one row per
position; sheet "settings": what was set and what the devices read back),
<name>.csv (the same rows, rewritten after every point, so a stopped run keeps
what it measured) and <name>.png (power vs. position; counts and voltage vs.
power).

Spacing: the ND filter's density is linear in position, so evenly spaced
positions (--spacing mm, the default) give powers evenly spaced on a LOG scale.
--spacing power picks positions whose powers are evenly spaced on a LINEAR
scale between the powers at --start and --stop, using the slope from
power_to_position.py (calibrated at 550 nm; only the slope is used, so the
laser's power doesn't matter — the measured powers are what's recorded anyway).

The positions run from --start to --stop in the order given; the default
90 -> 60 mm goes from low power to high, the gentle way for the detector.

The laser: with the mirror up and the ND stage at the range's high-power end,
the script turns emission on (if it's off; --power-level sets the level) and
watches the power meter: at least --warmup s (3), then until the power holds
steady (the last second's mean within --stable % of the second before), at
most --warmup-max s. That trace is the spreadsheet's "warmup" sheet. Emission
goes OFF at the end — also when the run stops early (an error, Ctrl+C). --leave-laser-on
skips that; --laser "" leaves the laser alone altogether.

A Script object on the canvas can run it too (scripts/canvas_run.py): its
inputs are these arguments, it shows the progress, and Stop ends the run after
the current point — with the usual cleanup (files saved, laser off, mirror back).
Its runs are the datasets of experiments/detector_sweep/, for a Graph to show.

Like calibrate_nd_filter.py, it goes through the instrument server (started if
it isn't running), needs the ND stage and the band homed, and the detection mirror's position known on the canvas. The mirror is
put back where it was at the end. The power is measured once, before the
detector pass: if the laser drifts in between, the two passes disagree by that
much.
"""

import argparse
import csv
import math
import os
import sys
import time
from datetime import datetime

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

from calibrate_nd_filter import (READ_PERIOD_S, Abort, average_power, connect,   # noqa: E402
                                 device_ok, move_stage, reading, set_mirror, wait_for)
import canvas_run                                                                 # noqa: E402
import power_to_position                                                          # noqa: E402

DATA_FOLDER = "detector_sweep"                  # experiments/detector_sweep/: one run per CSV
OUT_DIR = os.path.join(canvas_run.EXPERIMENTS, DATA_FOLDER)
BAND_TIMEOUT_S = 60.0
BAND_TOL_NM = 1.0
LASER_TIMEOUT_S = 30.0

COLUMNS = [                     # (key, header in the spreadsheet)
    ("index", "index"),
    ("target_mm", "target (mm)"),
    ("position_mm", "position (mm)"),
    ("power_mW", "power (mW)"),
    ("power_std_mW", "power std (mW)"),
    ("power_time", "power measured at"),
    ("count_rate", "count rate (counts/s)"),
    ("count_rate_std", "count rate std (counts/s)"),
    ("voltage_V", "voltage (V)"),
    ("voltage_std_V", "voltage std (V)"),
    ("overflows", "overflows (max)"),
    ("clipped", "clipped (max)"),
    ("detector_position_mm", "position, detector pass (mm)"),
    ("detector_time", "detector measured at"),
]


def now():
    return datetime.now().isoformat(sep=" ", timespec="seconds")


def mean_std(values):
    values = [v for v in values if v is not None]
    if not values:
        return None, None
    mean = sum(values) / len(values)
    return mean, (sum((v - mean) ** 2 for v in values) / max(1, len(values) - 1)) ** 0.5


def positions_for(start, stop, points, spacing):
    if spacing == "mm":
        return [start + (stop - start) * i / (points - 1) for i in range(points)]
    # Evenly spaced powers: P(x) ∝ 10 ** (SLOPE * x), so x = start + log10(P / P_start) / SLOPE
    slope = power_to_position.SLOPE
    ratio = 10 ** (slope * (stop - start))                  # P(stop) / P(start)
    return [start + math.log10(1 + (ratio - 1) * i / (points - 1)) / slope for i in range(points)]


def set_band(client, band, center, width):
    """Width first (keeps the center), then center (keeps the width); wait for both."""
    if reading(client, band, "homed") != 1:
        raise Abort(f"{band} (the excitation filters) isn't homed — home it on the canvas first")
    params = client.device(band)["params"]
    w_lo, w_hi = params["width"]["min"], params["width"]["max"]
    if not w_lo <= width <= w_hi:
        print(f"Note: the band can be {w_lo:g}–{w_hi:g} nm wide; {width:g} nm becomes the nearest.")
        width = max(w_lo, min(w_hi, width))
    for name, value in (("width", width), ("center", center)):
        client.set(band, name, value)
        time.sleep(0.3)                                     # let the set reach the stages

        def there():
            v = client.device(band)["values"]
            return (v.get("moving") == 0
                    and abs((v.get("actual_center") or 0) - center) <= BAND_TOL_NM
                    if name == "center" else v.get("moving") == 0)
        try:
            wait_for(there, BAND_TIMEOUT_S, f"the band's {name}")
        except Abort:
            v = client.device(band)["values"]
            raise Abort(f"the band didn't reach {center:g} nm × {width:g} nm (it's at "
                        f"{v.get('actual_center')} nm × {v.get('actual_width')} nm) — "
                        f"{client.device(band)['message'] or 'out of the filters range?'}")
    v = client.device(band)["values"]
    return v.get("actual_center"), v.get("actual_width")


def pda_average(client, pda, samples, gate):
    """Mean / std of `samples` fresh PDA45 readings. A reading is fresh when it
    differs from the one before — and the first change after arriving is
    thrown away, since that gate may have started before the stage stopped."""
    def latest():
        v = client.device(pda)["values"]
        return v.get("count_rate"), v.get("voltage"), v.get("overflows"), v.get("clipped")

    timeout = 3 * gate + 2.0
    got, last = [], latest()
    for k in range(samples + 1):
        new = wait_for(lambda: (lambda r: r if r != last else None)(latest()), timeout,
                       f"a fresh {pda} reading")
        last = new
        if k:
            got.append(new)
    rate, rate_std = mean_std([g[0] for g in got])
    volts, volts_std = mean_std([g[1] for g in got])
    if rate is None and volts is None:
        d = client.device(pda)["values"]
        raise Abort(f"{pda} gave neither counts nor voltage (tagger: {d.get('tagger_status')}, "
                    f"daq: {d.get('daq_status')})")
    worst = lambda i: max((g[i] for g in got if g[i] is not None), default=None)   # noqa: E731
    return {"count_rate": rate, "count_rate_std": rate_std, "voltage_V": volts,
            "voltage_std_V": volts_std, "overflows": worst(2), "clipped": worst(3)}


def laser_on(client, args):
    """Set the power level (if asked), turn emission on if it's off, and wait for it."""
    laser = args.laser
    if args.power_level is not None:
        client.set(laser, "power_level", args.power_level)
        print(f"Laser power level: {args.power_level:g} %")
    if reading(client, laser, "emission") == 1:
        print("Laser emission is already on")
        return
    if reading(client, laser, "interlock_ok") != 1:
        raise Abort(f"{laser}'s interlock is open ({reading(client, laser, 'interlock')}) — "
                    f"close it and Reset interlock on the canvas")
    print("Laser emission on…")
    client.send({"op": "action", "device": laser, "name": "emission_on"})
    try:
        wait_for(lambda: reading(client, laser, "emission") == 1, LASER_TIMEOUT_S, "emission on")
    except Abort:
        raise Abort(f"{laser} didn't turn on — {client.device(laser)['message'] or client.last_error or 'no reason given'}")
    print("  emission on")


def warm_up(client, args):
    """Watch the power meter until the power has settled: at least --warmup s,
    then until the mean of the last second is within --stable % of the second
    before it. Gives up after --warmup-max s (a warning, not a stop). Returns
    the trace, [(s, mW)], and how long it took."""
    print(f"Warm-up: watching the power (at least {args.warmup:g} s, until it holds within "
          f"{args.stable:g} %/s)")
    trace, t0, next_print = [], time.monotonic(), 0.0
    per_s = max(2, round(1.0 / READ_PERIOD_S))           # readings in one second
    while True:
        canvas_run.check_stop()
        time.sleep(READ_PERIOD_S)
        t = time.monotonic() - t0
        watts = reading(client, args.meter, "power")
        if watts is not None:
            trace.append((round(t, 3), watts * 1e3))
        if t >= next_print and trace:
            print(f"  {t:5.1f} s   {trace[-1][1]:10.5g} mW")
            next_print += 0.5
        if t >= args.warmup and len(trace) >= 2 * per_s:
            last = [mw for _, mw in trace[-per_s:]]
            before = [mw for _, mw in trace[-2 * per_s:-per_s]]
            now_mw, then_mw = sum(last) / len(last), sum(before) / len(before)
            change = abs(now_mw - then_mw) / now_mw * 100 if now_mw > 0 else float("inf")
            if change <= args.stable:
                print(f"  settled after {t:.1f} s: {now_mw:.5g} mW ({change:.2g} % in the last second)")
                return trace, t
        if t >= args.warmup_max:
            print(f"  WARNING: still changing after {t:.0f} s — going on anyway")
            return trace, t


def laser_off(client, laser):
    client.send({"op": "action", "device": laser, "name": "emission_off"})
    try:
        wait_for(lambda: reading(client, laser, "emission") == 0, LASER_TIMEOUT_S, "emission off")
        print("Laser emission off")
    except Abort:
        print(f"WARNING: couldn't confirm {laser}'s emission is off — check the laser!")


# --- Saving --------------------------------------------------------------------------
def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[k for k, _ in COLUMNS], extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(rows, settings, warmup, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    data = wb.active
    data.title = "data"
    data.append([h for _, h in COLUMNS])
    for row in rows:
        data.append([row.get(k) for k, _ in COLUMNS])
    for cell in data[1]:
        cell.font = Font(bold=True)
    for i, (_, header) in enumerate(COLUMNS, 1):
        data.column_dimensions[data.cell(1, i).column_letter].width = max(12, len(header) + 2)
    data.freeze_panes = "A2"

    info = wb.create_sheet("settings")
    for key, value in settings.items():
        info.append([key, value if isinstance(value, (int, float, str, type(None))) else str(value)])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 60
    for cell in info["A"]:
        cell.font = Font(bold=True)

    warm = wb.create_sheet("warmup")                     # the power while the laser came on
    warm.append(["time since emission on (s)", "power (mW)"])
    for row in warmup:
        warm.append(list(row))
    for cell in warm[1]:
        cell.font = Font(bold=True)
    warm.column_dimensions["A"].width = 28
    warm.column_dimensions["B"].width = 14
    wb.save(path)


def save_plot(rows, settings, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    powered = [r for r in rows if r.get("power_mW") is not None]
    both = [r for r in powered if r.get("count_rate") is not None or r.get("voltage_V") is not None]
    fig, (ax_p, ax_c) = plt.subplots(1, 2, figsize=(12, 4.6))

    ax_p.errorbar([r["position_mm"] for r in powered], [r["power_mW"] for r in powered],
                  yerr=[r["power_std_mW"] for r in powered], fmt="o-", ms=3, lw=1, capsize=2)
    ax_p.set_yscale("log")
    ax_p.set_xlabel("ND stage position (mm)")
    ax_p.set_ylabel("power (mW)")
    ax_p.set_title("power pass (mirror up)")
    ax_p.grid(True, which="both", alpha=0.25)

    ax_v = ax_c.twinx()
    counts = [r for r in both if r.get("count_rate") is not None]
    volts = [r for r in both if r.get("voltage_V") is not None]
    if counts:
        ax_c.errorbar([r["power_mW"] for r in counts], [r["count_rate"] for r in counts],
                      yerr=[r["count_rate_std"] for r in counts], fmt="o-", ms=3, lw=1,
                      capsize=2, color="tab:blue", label="count rate")
    if volts:
        ax_v.errorbar([r["power_mW"] for r in volts], [r["voltage_V"] for r in volts],
                      yerr=[r["voltage_std_V"] for r in volts], fmt="s--", ms=3, lw=1,
                      capsize=2, color="tab:orange", label="voltage")
    ax_c.set_xlabel("power (mW)")
    ax_c.set_ylabel("count rate (counts/s)", color="tab:blue")
    ax_v.set_ylabel("mean voltage (V)", color="tab:orange")
    ax_c.set_title("detector pass (mirror down)")
    ax_c.grid(True, alpha=0.25)

    fig.suptitle(f"PDA45 vs. power at {settings['wavelength_nm']:g} nm "
                 f"(band {settings.get('band_actual_center_nm')} nm × "
                 f"{settings.get('band_actual_width_nm')} nm)  ·  {settings['started']}")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


# --- The run -------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(description="PDA45 counts and voltage vs. power at one wavelength.")
    p.add_argument("--wavelength", type=float, required=True, help="nm: band center + meter correction")
    p.add_argument("--width", type=float, default=10.0, help="band width, nm (default 10)")
    p.add_argument("--start", type=float, default=90.0, help="first ND position, mm (default 90)")
    p.add_argument("--stop", type=float, default=60.0, help="last ND position, mm (default 60)")
    p.add_argument("--points", type=int, default=20, help="positions (default 20)")
    p.add_argument("--spacing", choices=("mm", "power"), default="mm",
                   help="mm = evenly spaced positions (log-spaced power, the default); "
                        "power = evenly spaced power")
    p.add_argument("--samples", type=int, default=5, help="power readings averaged per point (default 5)")
    p.add_argument("--pda-samples", type=int, default=5,
                   help="PDA45 readings averaged per point (default 5)")
    p.add_argument("--gate", type=float, default=None,
                   help="PDA45 gate per reading, s (default: leave it as it is)")
    p.add_argument("--settle", type=float, default=0.3, help="s to wait after each move (default 0.3)")
    p.add_argument("--name", default="", help="added to the file names, e.g. a sample name")
    p.add_argument("--stage", default="zaber", help="the ND stage's device name (default zaber)")
    p.add_argument("--meter", default="pm100d", help="the power meter's device name (default pm100d)")
    p.add_argument("--mirror", default="flipper", help="the detection flip mirror's device name")
    p.add_argument("--band", default="band", help="the excitation band's device name (default band)")
    p.add_argument("--pda", default="pda45", help="the detector's device name (default pda45)")
    p.add_argument("--laser", default="superk",
                   help='the laser\'s device name ("" = don\'t touch the laser; turn it on yourself)')
    p.add_argument("--power-level", type=float, default=None,
                   help="laser power setpoint, %% of max (default: leave it as it is)")
    p.add_argument("--warmup", type=float, default=3.0,
                   help="s to watch the power after emission on, at least (default 3)")
    p.add_argument("--warmup-max", type=float, default=30.0,
                   help="s to watch it at most before going on anyway (default 30)")
    p.add_argument("--stable", type=float, default=1.0,
                   help="settled = the last second's mean power within this %% of the second "
                        "before (default 1)")
    p.add_argument("--leave-laser-on", action="store_true", help="don't turn emission off at the end")
    p.add_argument("--home", action="store_true", help="home the ND stage first")
    job = canvas_run.start(p, data_folder=DATA_FOLDER, plot=("power_mW", "count_rate"))
    args = job.args
    if args.points < 2:
        p.error("--points must be at least 2")

    client = None
    try:
        client = connect()
        run(client, args)
        job.complete("experiment complete")
    except canvas_run.Stopped:
        print("\nStopped from the canvas, after the last point. The points so far are saved.")
        job.stopped("stopped after the last point; the points so far are saved")
    except Abort as exc:
        print(f"\nStopped: {exc}")
        job.fail(str(exc))
        sys.exit(1)
    except KeyboardInterrupt:
        print("\nStopped (Ctrl+C). The points so far are saved.")
        job.stopped("stopped with Ctrl+C; the points so far are saved")
        sys.exit(1)
    finally:
        if client is not None:
            client.enabled = False
            time.sleep(0.2)             # let queued requests go out


def run(client, args):
    """The measurement, with the laser turned off at the end whatever happens."""
    if args.laser:
        device_ok(client, args.laser)
    try:
        measure(client, args)
    finally:
        if args.laser and not args.leave_laser_on:
            laser_off(client, args.laser)


def measure(client, args):
    stage = device_ok(client, args.stage)
    device_ok(client, args.meter)
    device_ok(client, args.pda)
    device_ok(client, args.band)
    mirror = device_ok(client, args.mirror)

    if args.home:
        print("Homing the ND stage…")
        client.send({"op": "action", "device": args.stage, "name": "home"})
        time.sleep(0.5)
        wait_for(lambda: reading(client, args.stage, "homed") == 1
                 and reading(client, args.stage, "moving") == 0, 60.0, "homing")
    elif reading(client, args.stage, "homed") != 1:
        raise Abort(f"{args.stage} isn't homed — home it on the canvas, or run with --home")

    low, high = stage["params"]["target"]["min"], stage["params"]["target"]["max"]
    for name in ("start", "stop"):
        if not low <= getattr(args, name) <= high:
            raise Abort(f"--{name} {getattr(args, name):g} mm is outside the stage's {low:g}–{high:g} mm")
    positions = [round(x, 4) for x in positions_for(args.start, args.stop, args.points, args.spacing)]

    was_up = mirror["values"].get("is_up")
    if was_up is None:
        raise Abort(f"{args.mirror}'s position is unknown (no readback) — on the canvas, "
                    f"use It's at … or Find position first")

    canvas_run.stage("setting the band")
    print(f"Band: {args.wavelength:g} nm, {args.width:g} nm wide…")
    center, width = set_band(client, args.band, args.wavelength, args.width)
    print(f"  band at {center} nm × {width} nm")
    client.set(args.meter, "wavelength", args.wavelength)
    if args.gate is not None:
        client.set(args.pda, "gate", args.gate)
        time.sleep(0.5)
    gate = client.device(args.pda)["params"]["gate"]["value"] or 0.2

    os.makedirs(OUT_DIR, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    label = f"_{args.name}" if args.name else ""
    base = os.path.join(OUT_DIR, f"sweep_{args.wavelength:g}nm{label}_{stamp}")
    settings = {"started": now(), "wavelength_nm": args.wavelength, "width_nm": args.width,
                "band_actual_center_nm": center, "band_actual_width_nm": width,
                "start_mm": args.start, "stop_mm": args.stop, "points": args.points,
                "spacing": args.spacing, "power_samples": args.samples,
                "pda_samples": args.pda_samples, "pda_gate_s": gate, "settle_s": args.settle,
                "devices": f"stage {args.stage}, meter {args.meter}, mirror {args.mirror}, "
                           f"band {args.band}, detector {args.pda}, laser {args.laser or '-'}",
                "spacing_slope_dex_per_mm": power_to_position.SLOPE if args.spacing == "power" else None}
    warmup = []
    rows = [{"index": i, "target_mm": x} for i, x in enumerate(positions, 1)]
    canvas_run.dataset(base + ".csv")                    # this run's dataset (written after each point)
    total = 2 * len(rows)                                # a power pass, then a detector pass
    print(f"{args.points} points, {args.start:g} → {args.stop:g} mm ({args.spacing} spacing)")
    print(f"Writing {os.path.relpath(base, ROOT)}.xlsx / .csv / .png\n")

    try:
        set_mirror(client, args.mirror, up=True)
        print("Detection mirror up (beam -> power meter)")
        watch_at = min(positions)                        # the most power: the clearest signal
        move_stage(client, args.stage, watch_at)
        canvas_run.stage("laser warm-up")
        if args.laser:
            laser_on(client, args)
        warmup, took = warm_up(client, args)
        settings.update({"warmup_at_mm": watch_at, "warmup_s": round(took, 2),
                         "warmup_final_mW": warmup[-1][1] if warmup else None})
        if args.laser:
            laser = client.device(args.laser)
            settings.update({f"laser {k}": v for k, v in laser["values"].items()
                             if isinstance(v, (int, float, str))})
            settings["laser power_level"] = laser["params"]["power_level"]["value"]

        print("\nPower pass")
        canvas_run.progress(0, total, "power pass")
        for row in rows:
            canvas_run.check_stop()
            at = move_stage(client, args.stage, row["target_mm"])
            time.sleep(args.settle)
            watts, std = average_power(client, args.meter, args.samples)
            row.update(position_mm=at, power_mW=watts * 1e3, power_std_mW=std * 1e3, power_time=now())
            write_csv(rows, base + ".csv")
            print(f"  {row['index']:3d}/{args.points}  {at:8.3f} mm   {watts * 1e3:10.5g} mW  ± {std * 1e3:.2g}")
            canvas_run.progress(row["index"], total, "power pass")

        print("\nDetector pass — detection mirror down (beam -> sphere -> PDA45)")
        set_mirror(client, args.mirror, up=False)
        for row in rows:
            canvas_run.check_stop()
            at = move_stage(client, args.stage, row["target_mm"])
            time.sleep(args.settle)
            row.update(pda_average(client, args.pda, args.pda_samples, gate),
                       detector_position_mm=at, detector_time=now())
            write_csv(rows, base + ".csv")
            rate, volts = row["count_rate"], row["voltage_V"]
            flags = "".join([" OVERFLOW" if row["overflows"] else "", " CLIPPED" if row["clipped"] else ""])
            print(f"  {row['index']:3d}/{args.points}  {at:8.3f} mm   "
                  f"{'-' if rate is None else f'{rate:10.5g}'} counts/s   "
                  f"{'-' if volts is None else f'{volts:.5g}'} V{flags}")
            canvas_run.progress(len(rows) + row["index"], total, "detector pass")
    finally:
        settings["finished"] = now()
        done = [r for r in rows if "position_mm" in r]
        if done:
            write_csv(rows, base + ".csv")
            try:
                write_xlsx(rows, settings, warmup, base + ".xlsx")
                print(f"\nSpreadsheet: {os.path.relpath(base, ROOT)}.xlsx")
            except ImportError:
                print("\nNo openpyxl (pip install openpyxl) — the data is in the CSV only.")
            if len(done) >= 2:
                save_plot(rows, settings, base + ".png")
                print(f"Plot: {os.path.relpath(base, ROOT)}.png")
        current = reading(client, args.mirror, "is_up")
        if current is not None and current != was_up:
            set_mirror(client, args.mirror, up=bool(was_up))
            print(f"Detection mirror back {'up' if was_up else 'down'}")
    print("Done.")


if __name__ == "__main__":
    main()
